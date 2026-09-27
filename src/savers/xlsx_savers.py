"""Дочерний класс записи данных парсинга в формате XLSX."""

import shutil
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook

from core.base_classes import BaseSaver


class XLSXSaver(BaseSaver):
    """Сейвер для формата Excel (XLSX)."""

    def __init__(self, file_name: str):
        super().__init__(file_name)
        # Объявляем пути строго через self на уровне экземпляра класса
        self.file_path = (Path.cwd() / self.file_name).with_suffix(".xlsx")
        self.backup_path = (Path.cwd() / self.file_name).with_suffix(".xlsx.bak")
        self.is_file_init = False

    def save(self, data: List[Dict[str, Any]]) -> bool:
        if not data:
            return False

        # --- Делаем бэкап перед КАЖДОЙ записью чанка ---
        if self.file_path.exists():
            try:
                # Если прошлый бэкап остался, удаляем его
                self.backup_path.unlink(missing_ok=True)
                # Клонируем текущее состояние файла в бэкап
                shutil.copy2(self.file_path, self.backup_path)

                # Если это самый первый запуск НОВОЙ сессии — очищаем старый файл
                if not self.is_file_init:
                    self.file_path.unlink()
            except PermissionError as e:
                raise IOError(
                    f"❌ Не удалось сделать бэкап файла {self.file_name}. "
                    f"Убедитесь, что он закрыт в Excel перед запуском!"
                ) from e

        fieldnames = list(data[0].keys())

        try:
            if not self.file_path.exists():
                wb = Workbook()
                ws = wb.active
                ws.title = "Parsing Result"
                ws.append(fieldnames)  # Пишем шапку
            else:
                wb = load_workbook(self.file_path)
                ws = wb.active

            # Дописываем строки чанка
            for item in data:
                ws.append([item.get(key, "") for key in fieldnames])

            wb.save(self.file_path)
            wb.close()

            # --- УСПЕХ: Чанк записан, старый бэкап больше не нужен ---
            self.backup_path.unlink(missing_ok=True)
            self.is_file_init = True
            return True

        except Exception as e:
            # --- ОТКАТ: Запись чанка упала, восстанавливаем файл из бэкапа текущего шага ---
            if self.backup_path.exists():
                # Удаляем битый/недописанный файл
                self.file_path.unlink(missing_ok=True)
                # Возвращаем копию, сделанную прямо перед этим чанком
                shutil.copy2(self.backup_path, self.file_path)
                self.backup_path.unlink(missing_ok=True)

            raise IOError(f"Ошибка записи XLSX (данные восстановлены до состояния перед чанком):\n{e}") from e
